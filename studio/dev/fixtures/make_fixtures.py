# -*- coding: utf-8 -*-
"""파서 시험용 실제 문서를 만든다. openpyxl / python-pptx 는 시험 환경에만 있으면 된다."""
import io
import os
import sys
import zipfile

OUT = sys.argv[1] if len(sys.argv) > 1 else '.'
os.makedirs(OUT, exist_ok=True)

# ---- xlsx ------------------------------------------------------------
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "요구사항"
rows = [
    ["구분", "요구사항", "공수(M/D)", "난이도"],
    ["출석", "온라인/오프라인 통합 출석부", 9, 4],
    ["출석", "주차별 출석 현황 보기", 5, 3],
    ["성적", "출석 결과 성적부 반영", 6, 4],
    ["", "", "", ""],
    ["비고", "지난 학기 데이터 마이그레이션 필요", None, 5],
]
for r in rows:
    ws.append(r)
ws["F2"] = "떨어진 셀"          # 범위 계산 확인용
ws2 = wb.create_sheet("일정")
ws2.append(["단계", "시작", "종료"])
ws2.append(["개발", "2026-03-02", "2026-05-29"])
wb.create_sheet("빈시트")        # 건너뛰어야 한다
wb.save(os.path.join(OUT, "sample.xlsx"))

# ---- 숨긴 시트가 섞인 xlsx -------------------------------------------
#
# 실무 관리대장은 10장 중 9장이 숨김인 일이 흔하다(2026-10-07, LXP 대장).
# 전부 읽으면 **엑셀을 열어 봐도 없는 내용**이 WBS 로 나온다.
hb = Workbook()
hs = hb.active
hs.title = "보이는시트"
hs.append(["구분", "요구사항"])
hs.append(["출석", "통합 출석부"])

h1 = hb.create_sheet("개정이력")
h1.append(["Revision History"])
h1.append(["Date", "2026-01-01"])
h1.sheet_state = "hidden"

h2 = hb.create_sheet("지난설계")
h2.append(["홈 메인 → LMS 홈"])
h2.sheet_state = "veryHidden"

hb.save(os.path.join(OUT, "hidden_sheets.xlsx"))

# 전부 숨김인 문서. 거르면 "읽을 글자가 없습니다" 가 되어 멀쩡한 문서를
# 못 쓴다고 말하게 된다 — 그때는 거르지 않는 것이 맞다.
#
# 엑셀도 openpyxl 도 **마지막 남은 보이는 시트는 못 숨긴다.** 다른 도구가
# 만든 파일에서는 나오므로 workbook.xml 을 직접 손봐 그 꼴을 만든다.
import re

_tmp = os.path.join(OUT, "_tmp_allhidden.xlsx")
ab = Workbook()
a0 = ab.active
a0.title = "숨김1"
a0.append(["숨겨도 이것뿐"])
a1 = ab.create_sheet("숨김2")
a1.append(["이것도 숨김"])
a1.sheet_state = "hidden"
ab.save(_tmp)

_dst = os.path.join(OUT, "all_hidden.xlsx")
_zin = zipfile.ZipFile(_tmp)
with zipfile.ZipFile(_dst, "w", zipfile.ZIP_DEFLATED) as _zout:
    for _it in _zin.infolist():
        _data = _zin.read(_it.filename)
        if _it.filename == "xl/workbook.xml":
            _data = re.sub(r'(sheetId="1"[^>]*?)state="visible"', r'\1state="hidden"',
                           _data.decode("utf-8")).encode("utf-8")
        _zout.writestr(_it, _data)
_zin.close()
os.remove(_tmp)

# ---- pptx ------------------------------------------------------------
from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
slides = [
    ("A대 LXP 고도화", "킥오프 2026-02"),
    ("범위", "1. 출석 통합\n2. 성적부 연동\n3. 마이그레이션"),
    ("일정", "개발 3~5월 / 테스트 5~6월"),
]
for title, body in slides:
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    s.placeholders[1].text = body
# 글자 없는 슬라이드 — 건너뛰어야 한다
prs.slides.add_slide(prs.slide_layouts[6])
prs.save(os.path.join(OUT, "sample.pptx"))

# ---- docx (손으로 만든다. python-docx 가 없다) --------------------------
DOC_XML = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body>
<w:p><w:r><w:t>요구사항 정의서</w:t></w:r></w:p>
<w:p><w:r><w:t>1. 출석 </w:t></w:r><w:r><w:t>통합</w:t></w:r></w:p>
<w:p><w:r><w:t>온라인과 오프라인 출석을</w:t></w:r><w:r><w:br/><w:t>하나로 본다.</w:t></w:r></w:p>
<w:p/>
<w:p><w:r><w:t>2. 성적부</w:t></w:r><w:r><w:tab/><w:t>연동</w:t></w:r></w:p>
<w:tbl><w:tr><w:tc><w:p><w:r><w:t>항목</w:t></w:r></w:p></w:tc>
<w:tc><w:p><w:r><w:t>비고</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
</w:body></w:document>'''

CT = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>'''

RELS = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>'''

with zipfile.ZipFile(os.path.join(OUT, "sample.docx"), "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("[Content_Types].xml", CT)
    z.writestr("_rels/.rels", RELS)
    z.writestr("word/document.xml", DOC_XML)

# ---- 망가진 파일 -------------------------------------------------------
with open(os.path.join(OUT, "broken.xlsx"), "wb") as f:
    f.write(b"PK\x03\x04 this is not really a zip body " + b"\x00" * 40)

# zip 이긴 한데 office 문서가 아닌 것
with zipfile.ZipFile(os.path.join(OUT, "notoffice.xlsx"), "w") as z:
    z.writestr("readme.txt", "안녕")

print("만든 것:", sorted(os.listdir(OUT)))
